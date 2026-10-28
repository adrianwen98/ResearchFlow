import json
import pandas as pd
import pytest
from openpyxl import load_workbook
from web_agent.workspace import Workspace
from web_agent.persistence import atomic_json
from web_agent.business_excel import export_business_excel,verify_business_excel
from web_agent.completion import define_requirements,record_result,evaluate_requirements,review_binding
from web_agent.tools import build_local_tools
from test_graph import settings

def sourced(tmp_path):
    ws=Workspace(tmp_path,'business')
    ws.observe('https://example.com/spec','Model A capacity 5000 mAh; price -12.5 USD')
    ws.record('Specification','https://example.com/spec','Capacity','capacity 5000 mAh; price -12.5 USD',require_observation=True)
    return ws

def comparison(value=5000,**kwargs):
    return [{'name':'Compare','mode':'comparison','columns':['Capacity'], 'rows':[{'label':'Model A','cells':{'Capacity':{'value':value,'source_id':'1','quote':'capacity 5000 mAh',**kwargs}}}]}]

def test_dynamic_comparison_numeric_missing_and_binding(tmp_path):
    ws=sourced(tmp_path);binding=review_binding(ws)
    export_business_excel(ws,comparison())
    assert not verify_business_excel(ws) and review_binding(ws)!=binding
    with pd.ExcelFile(ws.path/'task_result.xlsx') as book:
        assert book.sheet_names==['Compare','Provenance','Metadata']
        assert pd.read_excel(book)['Capacity'][0]==5000
    define_requirements(ws,[{'description':'Create comparison workbook','kind':'artifact','artifact':'task_result.xlsx'}])
    record_result(ws,'r1',[])
    assert evaluate_requirements(ws)['requirements_supported']
    saved=review_binding(ws)
    book=load_workbook(ws.path/'task_result.xlsx');book['Compare']['B2']=999;book.save(ws.path/'task_result.xlsx');book.close()
    assert verify_business_excel(ws) and review_binding(ws)!=saved
    assert not evaluate_requirements(ws)['requirements_supported']

@pytest.mark.parametrize('spec',[comparison(9000),comparison(5000,source_id='fake'),comparison(5000,quote='fabricated'),comparison(None),[{'name':'../bad','mode':'table'}]])
def test_unsupported_cells_and_paths_rejected(tmp_path,spec):
    ws=sourced(tmp_path)
    with pytest.raises(ValueError):export_business_excel(ws,spec)
    assert not (ws.path/'task_result.xlsx').exists()

def test_table_and_calculation_preserve_negative_values_and_escape_only_strings(tmp_path,settings):
    ws=sourced(tmp_path);file='public_test.csv'
    (ws.path/file).write_text('group,value\nA,-1.5\nA,-2.5\n=unsafe,2\n',encoding='utf-8')
    atomic_json(ws.path/'downloads.json',{file:{'url':'https://example.com/data.csv'}})
    tools={t.name:t for t in build_local_tools(ws,settings)}
    calc=json.loads(tools['aggregate_downloaded_table'].invoke({'filename':file,'group_by':'group','numeric_column':'value','method':'sum'}))
    export_business_excel(ws,[{'name':'Data','mode':'table','filename':file},{'name':'Stats','mode':'calculation','calculation_id':calc['calculation_id']}])
    assert not verify_business_excel(ws)
    book=load_workbook(ws.path/'task_result.xlsx',read_only=True)
    assert book['Data']['B2'].value==-1.5 and book['Data']['A4'].value=="'=unsafe"
    assert {r[0]:r[1] for r in list(book['Stats'].values)[1:]}['A']==-4
    book.close()

def test_missing_business_value_has_explicit_reason(tmp_path):
    ws=sourced(tmp_path);export_business_excel(ws,comparison(None,reason='Official page does not provide it'))
    assert not verify_business_excel(ws)

def test_untrusted_column_headers_are_not_excel_formulas(tmp_path):
    ws=sourced(tmp_path);file='public_header.csv'
    (ws.path/file).write_text('=1+1,value\nx,-1\n',encoding='utf-8')
    atomic_json(ws.path/'downloads.json',{file:{'url':'https://example.com/data.csv'}})
    export_business_excel(ws,[{'name':'Data','mode':'table','filename':file}])
    book=load_workbook(ws.path/'task_result.xlsx',read_only=True)
    assert book['Data']['A1'].value=="'=1+1" and book['Data']['A1'].data_type!='f'
    book.close()
