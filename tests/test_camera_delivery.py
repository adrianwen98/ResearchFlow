"""Camera export regressions from saved public records; no model/network IO."""
import copy
import json
from pathlib import Path
import pytest
from openpyxl import load_workbook
from web_agent.workspace import Workspace
from web_agent.persistence import atomic_json,read_json
from web_agent.delivery_schema import normalize_entry
from web_agent.business_excel import (BusinessExcelValidationError,export_business_excel,
    inspect_business_draft,repair_business_excel,verify_business_excel)


def camera_workspace(root,key='camera'):
    case=json.loads((Path(__file__).parent/'fixtures/camera_delivery_case.json').read_text('utf-8'))
    ws=Workspace(root,key)
    atomic_json(ws.path/'sources.json',case['sources'])
    (ws.path/'observations').mkdir()
    for obs in case['observations']:atomic_json(ws.path/'observations'/f"{obs['id']}.json",obs)
    atomic_json(ws.path/'business_draft.json',case['draft'])
    return ws,case['draft']


def corrections(draft):
    spec=draft['sheets'][0];name=spec['name']
    fields=[{'sheet':name,'column':c,'field':{'type':'number','unit':u,'required':True}}
            for c,u in [('4K常规视频最高帧率','fps'),('最大码率','Mbps'),('标称续航','min')]]
    patches=[]
    for row in (2,3):
        entry=copy.deepcopy(spec['rows'][row-2]['cells']['内置存储'])
        entry['value']='The camera does not have built-in storage'
        patches.append({'sheet':name,'row':row,'column':'内置存储','entry':entry})
    return patches,fields


def test_actual_saved_camera_draft_repairs_in_one_batch_and_preserves_good_cells(tmp_path):
    ws,draft=camera_workspace(tmp_path)
    # Replay the failing input before patching; diagnostics must expose unit errors.
    inspected=inspect_business_draft(ws,draft['draft_id'])
    assert inspected['issues_total']==8
    assert any('field_patches' in i['message'] for i in inspected['issues'])
    bitrate=next(i for i in inspected['issues'] if i['column']=='最大码率')
    assert 'Max Video Bitrate' in bitrate['source_context']
    assert bitrate['field']['unit']=='count'
    original=copy.deepcopy(draft['sheets'][0])
    patches,fields=corrections(draft)
    repair_business_excel(ws,draft['draft_id'],patches,fields)
    assert verify_business_excel(ws)==[]
    meta=read_json(ws.path/'business_artifact.json')
    assert len(meta['provenance'])==32
    for p in meta['provenance']:
        if p['column'] not in {'内置存储','4K常规视频最高帧率','最大码率','标称续航'}:
            assert p['normalized_value']==original['rows'][p['row']-2]['cells'][p['column']]['value']
    with_book=load_workbook(ws.path/'task_result.xlsx')
    try:
        table=with_book[original['name']]
        assert table.freeze_panes=='B2' and table.tables
        headers=[c.value for c in table[1]]
        assert '最大码率 (Mbps)' in headers and '标称续航 (min)' in headers
        assert table.cell(2,headers.index('最大码率 (Mbps)')+1).value==130
    finally:with_book.close()
    assert not (ws.path/'business_draft.json').exists()
    assert not (ws.path/'usage.json').exists()  # no model budget touched
    assert all(p['entry']['value'] is not None for p in patches)


@pytest.mark.parametrize('value,raw,unit,input_unit,quote',[
    (60,'60fps','fps','fps','3840×2160@24/25/30/48/50/60fps'),
    (120,'120','fps','fps','3840×2160@100/120fps'),
    (130,'130 Mbps','Mbps','Mbps','Max Video Bitrate\n130 Mbps'),
    (130,'130000 kbps','Mbps','kbps','Bitrate 130000 kbps'),
    (166,'166 mins','min','min','Operating Time\n166 mins'),
    (120,2,'min','h','Operating Time 2 hours'),
    (5,'5','mAh','mAh','Capacity 5 mAh'),
])
def test_proven_scalar_and_dimension_conversion(value,raw,unit,input_unit,quote):
    actual,audit=normalize_entry({'value':value,'raw_value':raw,'input_unit':input_unit,'quote':quote},
                               {'type':'number','unit':unit,'required':True},None)
    assert actual==value and type(audit['raw_value']) in (float,int)
    assert audit['input_unit']==input_unit


@pytest.mark.parametrize('raw,input_unit,unit,quote',[
    ('60fps','fps','count','60fps'),
    ('130 MBps','Mbps','Mbps','130 MBps'),
    (130,'Mbps','Mbps','130 MB/s'),
    (130,'Mbps','Mbps','130 pixels and 50 Mbps'),
    ('130 Mbps extra','Mbps','Mbps','130 Mbps extra'),
    (130,'fps','fps','130fpsx'),
    ('9007199254740993','count','count','9007199254740993 count'),
])
def test_bad_units_unbound_values_prose_and_precision_still_rejected(raw,input_unit,unit,quote):
    with pytest.raises(ValueError):
        normalize_entry({'value':130,'raw_value':raw,'input_unit':input_unit,'quote':quote},
                        {'type':'number','unit':unit,'required':True},None)


def test_field_repairs_are_task_bound_transactional_and_cannot_drop_required(tmp_path):
    ws,draft=camera_workspace(tmp_path);path=ws.path/'business_draft.json';before=path.read_bytes()
    patches,fields=corrections(draft)
    invalid=copy.deepcopy(fields);invalid[0]['field']['required']=False
    with pytest.raises(ValueError,match='必需'):repair_business_excel(ws,draft['draft_id'],patches,invalid)
    assert path.read_bytes()==before
    with pytest.raises(ValueError,match='第1个.*row=None'):
        repair_business_excel(ws,draft['draft_id'],[{'sheet':fields[0]['sheet'],'column':'内置存储','entry':{'value':None,'reason':'unknown'}}])
    assert path.read_bytes()==before
    other=Workspace(tmp_path,'other')
    with pytest.raises(ValueError):repair_business_excel(other,draft['draft_id'],patches,fields)
    with pytest.raises(ValueError):repair_business_excel(ws,'0'*24,patches,fields)
    assert path.read_bytes()==before


def test_field_only_repair_rechecks_every_cell_and_retains_text_errors(tmp_path):
    ws,draft=camera_workspace(tmp_path);_,fields=corrections(draft)
    with pytest.raises(BusinessExcelValidationError) as error:
        repair_business_excel(ws,draft['draft_id'],[],fields)
    assert error.value.details['issues_total']==2
    assert {i['column'] for i in error.value.issues}=={'内置存储'}
    assert not (ws.path/'task_result.xlsx').exists()
    newer=read_json(ws.path/'business_draft.json')
    assert newer['draft_id']!=draft['draft_id']
    with pytest.raises(ValueError):repair_business_excel(ws,draft['draft_id'],[],fields)
