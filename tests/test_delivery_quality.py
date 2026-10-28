import hashlib
import json
import pytest
import pymupdf
from openpyxl import load_workbook
from web_agent.workspace import Workspace
from web_agent.persistence import atomic_json,read_json
from web_agent.business_excel import export_business_excel,BusinessExcelValidationError
from web_agent.report_export import render_pdf,safe_html
from web_agent.completion import evaluate_requirements,define_requirements,record_result

@pytest.fixture
def ws(tmp_path):
    w=Workspace(tmp_path,'quality')
    text='Screen 160 mm. Battery 5 Ah. Processor Alpha 8. Buy from $899. United States. launch price USD 799.'
    w.observe('https://example.com/us/shop',text)
    w.record('Official','https://example.com/us/shop','Specs',text,require_observation=True)
    atomic_json(w.path/'task.json',{'task':'Quality test','report_format':'both'})
    return w

def sheet(col,entry,fields=None):
    return [{'name':'Compare','mode':'comparison','columns':[col], 'fields':fields or {},'rows':[{'label':'Item A','cells':{col:entry}}]}]

def test_units_typed_values_style_and_audit(ws):
    export_business_excel(ws,sheet('电池容量',{'value':5000,'raw_value':5,'input_unit':'Ah','source_id':'1','quote':'Battery 5 Ah'}))
    meta=read_json(ws.path/'business_artifact.json')
    assert meta['sheets'][0]['rows']==[['Item A',5000]]
    assert meta['provenance'][0]['rule']=='Ah -> mAh'
    with_book=load_workbook(ws.path/'task_result.xlsx')
    try:
        table=with_book['Compare']
        assert table['B2'].data_type=='n' and table.freeze_panes=='B2'
        assert table.tables and table['B2'].alignment.wrap_text and table.conditional_formatting
    finally:with_book.close()

@pytest.mark.parametrize('col,entry',[
    ('屏幕尺寸',{'value':'6.3 inch OLED display and other unrelated advertising'}),
    ('电池容量',{'value':True,'quote':'Battery 5 Ah'}),
    ('电池容量',{'value':5000,'raw_value':5,'input_unit':'g','quote':'Battery 5 Ah'}),
    ('电池容量',{'value':6000,'raw_value':5,'input_unit':'Ah','quote':'Battery 5 Ah'}),
    ('处理器',{'value':'4.47GHz, 3.5GHz'}),
    ('Description',{'value':'x'*81}),
    ('电池容量',{'value':None,'status':'not_disclosed','reason':'not found','quote':'Battery 5 Ah'}),
    ('价格',{'value':899,'quote':'Buy from $899'}),
])
def test_bad_scalar_rejected_without_workbook(ws,col,entry):
    with pytest.raises(BusinessExcelValidationError):
        export_business_excel(ws,sheet(col,{'source_id':'1','quote':ws.sources()[0]['quote'],**entry}))
    assert not (ws.path/'task_result.xlsx').exists()

def price_entry(ws,**overrides):
    p={'currency':'USD','market':'US','type':'starting','queried_at':ws.sources()[0]['accessed_at'][:10],**overrides}
    return {'value':899,'quote':'Buy from $899','source_id':'1','price':p}

@pytest.mark.parametrize('overrides',[{'currency':'SGD'},{'market':'CN'},{'type':'launch'},{'queried_at':'2000-01-01'}])
def test_price_context_cannot_be_faked(ws,overrides):
    with pytest.raises(BusinessExcelValidationError):export_business_excel(ws,sheet('价格',price_entry(ws,**overrides)))

def test_price_metadata_and_missing_block_completion(ws):
    export_business_excel(ws,sheet('价格',price_entry(ws)))
    meta=read_json(ws.path/'business_artifact.json')
    assert meta['sheets'][0]['rows'][0][1:]==[899,'USD','US','起售价',ws.sources()[0]['accessed_at'][:10]]
    export_business_excel(ws,sheet('电池容量',{'value':None,'reason':'No verified capacity'}))
    ws.export_excel();ws.export_report('# Result\nSome observed facts.')
    assert 'No verified capacity' in (ws.path/'research_report.md').read_text(encoding='utf-8')
    define_requirements(ws,[{'description':'Business workbook','kind':'artifact','artifact':'task_result.xlsx'}])
    record_result(ws,'r1',[])
    assert not evaluate_requirements(ws)['requirements_supported']

def test_pdf_same_source_embedded_cjk_links_and_tamper(ws):
    ws.export_excel();ws.export_report('# 中文报告\n\n## 执行摘要\n中文正常显示。\n\n| 项目 | 数值 |\n|---|---|\n| 参数 | 6.3 |')
    meta=read_json(ws.path/'report_meta.json')
    assert meta['pdf_source_sha256']==hashlib.sha256((ws.path/'research_report.md').read_bytes()).hexdigest()
    assert ws.verify_delivery()['artifacts_valid']
    with pymupdf.open(ws.path/'research_report.pdf') as doc:
        text=''.join(p.get_text() for p in doc)
        assert '中文报告' in text and '6.3' in text and '执行摘要' in text
        assert any(p.get_links() for p in doc)
        assert all(p.get_text().strip() for p in doc)
    (ws.path/'research_report.pdf').write_bytes(b'not a PDF')
    assert not ws.verify_delivery()['artifacts_valid']

def test_no_external_resources_or_unsafe_links(tmp_path):
    text='# Safe\n![image](file:///secret.txt)\n<script>alert(1)</script>\n[bad](javascript:alert(1))\n[local](http://127.0.0.1/private)'
    html=safe_html(text)
    assert '<img' not in html and '<script>' not in html and 'href="http://127' not in html
    render_pdf(text,tmp_path/'safe.pdf')
    with pymupdf.open(tmp_path/'safe.pdf') as doc:assert not any(p.get_links() for p in doc)

def test_long_table_all_rows_and_no_empty_pages(tmp_path):
    md='# 多页表格\n\n| ID | 说明 |\n|---|---|\n'+'\n'.join(f'| ROW{i:03} | 第{i}项，记录连续原文与规范化结果，检查长表格的分页。 |' for i in range(100))
    render_pdf(md,tmp_path/'long.pdf')
    with pymupdf.open(tmp_path/'long.pdf') as doc:
        assert 2<len(doc)<30
        text=''.join(p.get_text() for p in doc)
        assert all(f'ROW{i:03}' in text for i in range(100))
        assert all(len(p.get_text().strip())>120 for p in doc)

@pytest.mark.parametrize('fmt,expected',[('markdown',False),('pdf',True),('both',True)])
def test_persisted_format_and_legacy_default(ws,fmt,expected):
    atomic_json(ws.path/'task.json',{'task':'Format test','report_format':fmt})
    ws.export_excel();ws.export_report('# Format\nSaved content')
    reopened=Workspace(ws.path.parent,ws.task_id)
    assert (reopened.path/'research_report.pdf').exists()==expected
    names=[p.split('\\')[-1].split('/')[-1] for p in reopened.list_artifacts()]
    assert ('research_report.md' in names)==(fmt!='pdf')
    assert reopened.verify_delivery()['artifacts_valid']
