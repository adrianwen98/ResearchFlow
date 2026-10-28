"""Evidence-backed requirements with version-bound human or narrow objective acceptance."""
import hashlib
import json
import re
from .persistence import atomic_json,read_json,read_trace
from .memory import memory_for
from .status import completion_outcome


def excerpt_in_report(excerpt,report):
    # Preserve words, numbers and punctuation; ignore only whitespace and common Markdown marks.
    normalize=lambda text:re.sub(r'\s+','',text.replace('**','').replace('`',''))
    return bool(excerpt.strip()) and normalize(excerpt) in normalize(report)


def define_requirements(workspace,items,*,origin='model'):
    if not 1<=len(items)<=16:raise ValueError('需求清单需要1–16项')
    path=workspace.path/'requirements.json'
    existing=read_json(path,[])
    for item in items:
        description=str(item.get('description','')).strip()
        kind=item.get('kind','research')
        if not 3<=len(description)<=600 or kind not in {'research','artifact','calculation'}:
            raise ValueError('每项需要明确description，kind为research/artifact/calculation')
        artifact = item.get('artifact') or ''

        if kind == 'artifact' and not artifact:
            desc = description.lower()

            if 'task_result.xlsx' in desc or any(
                    x in desc for x in ('对比表', '统计表', '数据表', 'excel')
            ):
                artifact = 'task_result.xlsx'

            elif 'research_sources.xlsx' in desc or '来源证据表' in desc:
                artifact = 'research_sources.xlsx'

            elif 'research_report.pdf' in desc or 'pdf报告' in desc:
                artifact = 'research_report.pdf'

            elif any(
                    x in desc for x in ('研究报告', '中文报告', 'markdown', '报告')
            ):
                artifact = 'research_report.md'

        if kind == 'artifact' and artifact not in {
            'research_report.md',
            'research_report.pdf',
            'research_sources.xlsx',
            'task_result.xlsx'
        }:
            raise ValueError(
                '无法确定交付文件。artifact必须指定为'
                'research_report.md、research_report.pdf、'
                'research_sources.xlsx或task_result.xlsx'
            )

        if any(row['description']==description for row in existing):continue
        if len(existing)>=16:raise ValueError('需求最多16项；不能通过替换删除旧需求')
        existing.append({'id':'r'+str(len(existing)+1),'description':description,'kind':kind,'artifact':artifact,
                         'origin':origin,'result':{'status':'pending'},
                         # Only user-supplied deterministic contracts may enable
                         # automatic completion. A model cannot certify itself.
                         'machine_check':item.get('machine_check') if origin=='user' else None})
    atomic_json(path,existing)
    workspace.log('requirements_defined','',count=len(existing),origin=origin)
    return existing


def record_result(workspace,requirement_id,source_ids,report_excerpt='',calculation_id='',status='satisfied',reason='',report_line_ids=None,report_sha256='',fact_refs=None):
    rows=read_json(workspace.path/'requirements.json',[])
    row=next((r for r in rows if r['id']==requirement_id),None)
    if row is None:raise ValueError('需求ID不存在，请先读取/定义需求')
    if status not in {'satisfied','missing','not_disclosed'}:raise ValueError('status只能satisfied、missing或not_disclosed')
    if not isinstance(source_ids,list):source_ids=[]
    auto_sources=False
    if status=='satisfied' and row['kind']=='calculation' and not source_ids:
        calculation=read_json(workspace.path/'calculations.json',{}).get(calculation_id,{})
        if isinstance(calculation,dict):
            source_id=calculation.get('source_id')
            if isinstance(source_id,str) and source_id:
                source_ids=[source_id]
            elif calculation.get('observation_id'):
                source_ids=[s['id'] for s in workspace.sources()
                            if s.get('observation_id')==calculation['observation_id']]
            auto_sources=bool(source_ids)
            if auto_sources:
                workspace.log('requirement_source_autofill', requirement_id,
                              calculation_id=calculation_id, source_ids=source_ids)
    excerpts=None
    if report_line_ids is not None:
        if status not in {'satisfied','not_disclosed'} or report_excerpt:raise ValueError('报告行引用仅用于satisfied/not_disclosed，不能同时填写report_excerpt')
        from .delivery_inspection import report_lines
        excerpts,report_sha256=report_lines(workspace,report_line_ids,report_sha256)
        report_excerpt='\n'.join(excerpts)
    known={source['id'] for source in workspace.sources()}
    if set(source_ids)-known:raise ValueError('引用了不存在的来源ID')
    if status in {'satisfied','not_disclosed'} and row['kind']!='artifact' and (not source_ids or not report_excerpt.strip()):
        if not source_ids:
            raise ValueError('该需求还没有已登记来源：请先调用 record_source 或 '
                             'record_source_from_observation 登记真实来源，再用 source_ids 引用；'
                             '不要重复提交相同参数')
        raise ValueError('该需求缺少最终报告中的连续原文：请先用 inspect_delivery 获取 '
                         'report_line_ids 和 report_sha256，或复制报告中真实存在的连续原文')
    report_path=workspace.path/'research_report.md'
    if status in {'satisfied','not_disclosed'} and row['kind']!='artifact' and report_path.exists() and not all(excerpt_in_report(e,report_path.read_text(encoding='utf-8')) for e in (excerpts or [report_excerpt])):
        from .delivery_inspection import DeliveryInputError,saved_report
        import difflib
        text,digest=saved_report(workspace)
        candidates=sorted([(difflib.SequenceMatcher(None,report_excerpt,line).ratio(),i+1,line) for i,line in enumerate(text.splitlines()) if line.strip()],reverse=True)[:3]
        raise DeliveryInputError('需求摘录与已导出报告不匹配；用inspect_delivery读取报告行，再用report_line_ids与report_sha256引用，不能改写或拼接原文',
              {'kind':'report_excerpt','report_sha256':digest,'next_tool':'inspect_delivery',
               'candidate_lines':[{'line_id':i,'text':line[:600]} for _,i,line in candidates]})
    if row['kind']=='calculation' and status=='satisfied':
        if calculation_id not in read_json(workspace.path/'calculations.json',{}):raise ValueError('计算结果ID不存在')
    # Explicit business-cell references survive later report-only updates.
    # They are an audit contract, not a model-certified assertion of coverage.
    from .requirement_fact_links import normalize_fact_refs
    if fact_refs is None:
        fact_refs=row.get('result',{}).get('fact_refs')
    if fact_refs is not None:
        if row['kind']!='research':raise ValueError('fact_refs仅适用于research需求')
        fact_refs=normalize_fact_refs(fact_refs)
    row['result']={'status':status,'source_ids':source_ids,'report_excerpt':report_excerpt[:2000],
                   'calculation_id':calculation_id,'reason':reason[:600]}
    if auto_sources:row['result']['auto_source_ids']=True
    if fact_refs is not None:row['result']['fact_refs']=fact_refs
    if excerpts is not None:row['result'].update(report_excerpts=excerpts,report_sha256=report_sha256,report_line_ids=sorted(report_line_ids))
    atomic_json(workspace.path/'requirements.json',rows)
    return row


def evaluate_requirements(workspace):
    rows=read_json(workspace.path/'requirements.json',[])
    report_path=workspace.path/'research_report.md'
    report=report_path.read_text(encoding='utf-8') if report_path.exists() else ''
    sources={row['id'] for row in workspace.sources()}
    calculations=read_json(workspace.path/'calculations.json',{})
    checks=[]
    for row in rows:
        result=row.get('result',{})
        status=result.get('status')
        valid=status=='satisfied'
        if row['kind']=='artifact':
            valid=valid and (workspace.path/row['artifact']).is_file()
            if row['artifact']=='task_result.xlsx':
                from .business_excel import verify_business_excel
                valid=valid and (workspace.path/'business_artifact.json').is_file() and not verify_business_excel(workspace)
        else:
            valid=status in {'satisfied','not_disclosed'} and bool(result.get('source_ids')) and set(result.get('source_ids',[]))<=sources
            parts=result.get('report_excerpts')
            valid=valid and (all(isinstance(e,str) and excerpt_in_report(e,report) for e in parts) if isinstance(parts,list) and parts else excerpt_in_report(result.get('report_excerpt',''),report))
            if result.get('report_sha256'):valid=valid and report_path.is_file() and result['report_sha256']==hashlib.sha256(report_path.read_bytes()).hexdigest()
        if row['kind']=='calculation':valid=valid and status=='satisfied' and result.get('calculation_id') in calculations
        reason=result.get('reason','')
        if not valid and not reason:reason='尚未验证' if result.get('status','pending')=='pending' else '文件、来源、计算ID或报告摘录不再满足检查，请重新补齐。'
        checks.append({'id':row['id'],'description':row['description'],'passed':bool(valid),
                       'status':result.get('status','pending'),'reason':reason})
    business=read_json(workspace.path/'business_artifact.json',{})
    candidates=[p for p in business.get('provenance',[]) if (p.get('required') and (p.get('status') in {'unverified','not_disclosed','review_required'} or p.get('price_status')=='unverified'))
                or (business.get('structured_review') and p.get('status')=='review_required')]
    # Field presence and factual certainty are separate. A disclosed unknown in
    # open-ended research goes to human review, never to automatic completion.
    # Legacy strict exports and explicit user research/calculation contracts stay strict.
    strict_contract=any(r.get('origin')=='user' and r.get('kind') in {'research','calculation'} for r in rows)
    warnings=[];blockers=[]
    report_bound=False
    if business.get('review_disclosed_missing') and not strict_contract:
        from .report_export import verify_report_formats
        report_bound=report_path.is_file() and read_json(workspace.path/'report_meta.json',{}).get('sha256')==hashlib.sha256(report_path.read_bytes()).hexdigest() and not verify_report_formats(workspace)
    for p in candidates:
        label=p['label']+' / '+p['column']
        disclosed=(p.get('status') in {'unverified','not_disclosed','review_required'} and p.get('note','').strip()
            and any(p['note'] in note and excerpt_in_report(note,report) for note in business.get('quality_notes',[])))
        if report_bound and disclosed and (p.get('price_status')!='unverified' or business.get('structured_review')):warnings.append(label+'：'+p['note'])
        else:blockers.append(label)
    # A model's "satisfied" is a claim that references were filed, not an
    # independent factual certification. Preserve the established human
    # acceptance policy, but expose source-scoped business uncertainties to
    # the reviewer instead of labelling the research itself verified.
    issues_by_source={}
    for p in candidates:
        sid=p.get('source_id')
        if isinstance(sid,str) and sid:
            label=str(p.get('label',''))+' / '+str(p.get('column',''))
            if label not in issues_by_source.setdefault(sid,[]):
                issues_by_source[sid].append(label)
    for requirement,check in zip(rows,checks):
        if requirement.get('kind')!='research':
            check['fact_status']='not_applicable'
            check['source_quality_issues']=[]
            continue
        referenced=requirement.get('result',{}).get('source_ids',[])
        related=[]
        if isinstance(referenced,list):
            for sid in referenced:
                if isinstance(sid,str):
                    for label in issues_by_source.get(sid,[]):
                        if label not in related:related.append(label)
        check['source_quality_issues']=related[:20]
        # No fuzzy natural-language matching: an uncertain fact in a cited
        # source is a review signal, NOT proof that this requirement is wrong.
        check['fact_status']=('unmet' if not check['passed'] else
                              'review_required' if related else
                              'not_disclosed' if requirement.get('result',{}).get('status')=='not_disclosed'
                              else 'not_independently_verified')
        # Only explicit exact cell IDs can bind a requirement to its business
        # facts. A common source ID is not enough to establish that binding.
        fact_refs=requirement.get('result',{}).get('fact_refs')
        if fact_refs is not None:
            from .requirement_fact_links import audit_fact_refs
            audit=audit_fact_refs(workspace,fact_refs,referenced,business)
            check['fact_links']=audit['checks']
            check['fact_link_issues']=audit['issues']
            if not audit['all_verified']:
                check['passed']=False
                check['fact_status']='unmet'
                check['reason']='显式关联的业务字段尚未全部核实：'+'；'.join(audit['issues'][:3])
            elif check['passed']:
                # Verified cell provenance is not proof that every element of
                # the natural-language requirement was covered.
                check['fact_status']='linked_cells_verified'
    return {'requirements':checks,'quality_blockers':blockers,'quality_warnings':warnings,
            'requirements_supported':bool(checks) and all(r['passed'] for r in checks) and not blockers,
            'missing_requirements':[r['id'] for r in checks if not r['passed']]}



def automatic_checks(workspace):
    """Narrow objective contracts; free-form research always needs judgment."""
    if evaluate_requirements(workspace).get('quality_warnings'):return False,'交付包含已披露的未核实信息，请人工核对并决定是否接受这些限制。'
    task=read_json(workspace.path/'task.json',{})
    if task.get('review_mode','auto')=='manual' or re.search(r'人工验收|人工确认|手动确认|manual acceptance|human review',task.get('task',''),re.I):return False, '任务明确指定人工确认。'
    rows=read_json(workspace.path/'requirements.json',[])
    if not rows:return False, '没有完整的逐项需求，不能自动确认。'
    if any(e.get('event') in {'approved_action_completed','human_decision'} for e in read_trace(workspace.path/'trace.jsonl')):
        return False, '任务包含人工处理或外部操作，需要确认最终结果。'
    sources={s['id']:s for s in workspace.sources()}
    report=(workspace.path/'research_report.md').read_text(encoding='utf-8') if (workspace.path/'research_report.md').exists() else ''
    for row in rows:
        check=row.get('machine_check')
        if row.get('origin')!='user' or not isinstance(check,dict):
            return False, '自由文本研究或模型拆分的需求需要用户判断；结构校验不等于事实正确。'
        try:
            kind=check.get('type')
            if kind=='source_quote' and row['kind']=='research':
                expected=check['quote'];url=check['url']
                if not isinstance(expected,str) or not expected.strip():return False,'自动检查缺少明确原文。'
                candidates=[sources[sid] for sid in row['result']['source_ids'] if sources[sid]['url']==url]
                if not candidates or not excerpt_in_report(expected,report):return False,'指定来源原文尚未完整出现在报告。'
                workspace.evidence_for(url,expected,candidates[0].get('observation_id'))
            elif kind=='calculation' and row['kind']=='calculation':
                from .tabular import aggregate_table
                calc=read_json(workspace.path/'calculations.json',{})[row['result']['calculation_id']]
                if any(calc.get(k,'')!=check.get(k,'') for k in ('group_by','numeric_column','method')):
                    return False,'计算参数与用户指定的客观要求不一致。'
                recomputed=json.loads(aggregate_table(workspace.path,calc['filename'],calc['group_by'],calc['numeric_column'],calc['method']))
                if recomputed!=calc['result'] or not excerpt_in_report(json.dumps(recomputed,ensure_ascii=False),report):
                    return False,'计算未能从原始数据复算或报告缺少完整统计原文。'
            elif kind=='artifact' and row['kind']=='artifact':
                if check.get('filename')!=row['artifact']:return False,'交付文件与指定文件不一致。'
                if row['artifact']=='task_result.xlsx':
                    from openpyxl import load_workbook
                    book=load_workbook(workspace.path/row['artifact'],read_only=True,data_only=True)
                    try:
                        sheet=book[check['sheet']]
                        columns=[str(v) for v in next(sheet.values)]
                        if not isinstance(check['columns'],list) or not check['columns'] or columns!=check['columns'] or (sheet.max_row or 0)-1 < max(1,int(check.get('min_rows',1))):
                            return False,'工作表结构或行数不满足指定要求。'
                    finally:book.close()
                elif row['artifact']=='research_report.md':
                    if not isinstance(check.get('contains'),list) or not check['contains'] or not all(isinstance(text,str) and excerpt_in_report(text,report) for text in check['contains']):
                        return False,'报告缺少指定的客观原文。'
                else:
                    # Source workbook still checked structurally by verify_delivery.
                    if row['artifact']!='research_sources.xlsx':return False,'不支持的自动文件检查。'
            else:return False,'需求含主观判断或未支持的自动检查。'
        except (ValueError,KeyError,TypeError,OSError,StopIteration):
            return False,'客观验收参数、来源或原始文件未通过检查。'
    return True, '全部用户指定的客观检查通过，已自动完成；不代表人工验收。'


def reconcile_final_answer(answer, status, requirements):
    """Remove a stale 'partial/missing' line after host artifact reconciliation."""
    if status not in {'needs_review','completed'} or requirements.get('missing_requirements'):
        return answer
    text = answer or ''
    if not text:
        return '任务已完成。' if status == 'completed' else '交付已生成，等待你核对。'
    if not text.startswith('已生成部分交付，尚未满足需求：'):
        return text
    remainder = text.split('\n', 1)[1].strip() if '\n' in text else ''
    prefix = '任务已完成。' if status == 'completed' else '交付已生成，等待你核对。'
    return prefix + ('\n' + remainder if remainder else '')


def completion_decision(workspace):
    artifacts=workspace.verify_delivery();requirements=evaluate_requirements(workspace)
    if not artifacts['artifacts_valid'] or (requirements['requirements'] and not requirements['requirements_supported']):
        return {'status':'partial' if workspace.sources() else 'incomplete','reason':('必要参数仍未核实：'+'；'.join(requirements['quality_blockers'])+'。已保留缺失说明，不能标记全部完成。') if requirements['quality_blockers'] else '存在文件/来源问题或未满足的必要需求，不能标记全部完成。'}
    automatic,reason=automatic_checks(workspace)
    return {'status':'completed' if automatic else 'needs_review','reason':reason}


def automatic_receipt(workspace):
    decision=completion_decision(workspace)
    if decision['status']!='completed':raise ValueError('自动验收条件不满足')
    receipt={'binding':review_binding(workspace),'verification_mode':'automatic_objective_checks',
             'requirement_ids':[r['id'] for r in read_json(workspace.path/'requirements.json',[])],
             'note':decision['reason']}
    atomic_json(workspace.path/'acceptance.json',receipt)
    workspace.log('delivery_auto_completed',decision['reason'])
    return receipt


def review_binding(workspace):
    names=['research_report.md','research_sources.xlsx','sources.json','requirements.json']
    if (workspace.path/'task_result.xlsx').exists() or (workspace.path/'business_artifact.json').exists():names+=['task_result.xlsx','business_artifact.json']
    if (workspace.path/'research_report.pdf').exists():names+=['research_report.pdf','report_meta.json']
    paths=[workspace.path/name for name in names]
    digest=hashlib.sha256()
    for path in paths:
        digest.update(path.name.encode());digest.update(path.read_bytes() if path.exists() else b'')
    return digest.hexdigest()


def accept_delivery(workspace,requirement_ids,binding,note='',*,accept_limitations=False):
    lock=workspace.path/'.running'
    try:handle=lock.open('x')
    except FileExistsError as exc:raise ValueError('任务正在运行，不能验收') from exc
    handle.close()
    try:return _accept_delivery(workspace,requirement_ids,binding,note,accept_limitations=accept_limitations)
    finally:lock.unlink(missing_ok=True)


def _accept_delivery(workspace,requirement_ids,binding,note,*,accept_limitations=False):
    if read_json(workspace.path/'demo_meta.json',{}).get('read_only'):raise ValueError('脱敏演示副本只读；不能作为新任务验收')
    if binding!=review_binding(workspace):raise ValueError('交付内容已变化，请重新打开任务核对后验收')
    artifacts=workspace.verify_delivery()
    if not artifacts['artifacts_valid']:raise ValueError('文件/来源核验未通过，不能验收：'+','.join(artifacts['issues']))
    rows=read_json(workspace.path/'requirements.json',[])
    if not rows or set(requirement_ids)!={row['id'] for row in rows}:
        raise ValueError('必须明确确认当前全部需求ID')
    requirements=evaluate_requirements(workspace)
    if not requirements['requirements_supported']:
        raise ValueError('存在未满足的需求，不能标记完成：'+','.join(requirements['missing_requirements']))
    if requirements.get('quality_warnings') and accept_limitations is not True:raise ValueError('请明确确认已阅读并接受本次交付的未核实信息；不会将未知参数标记为已核实')
    result=read_json(workspace.path/'result.json',{})
    if not result:raise ValueError('任务没有执行结果')
    if result.get('status')!='needs_review':raise ValueError('当前状态不能验收；请先重新核验交付，进入结果待确认状态')
    receipt={'binding':binding,'requirement_ids':requirement_ids,'note':note[:1000],'verification_mode':'human_acceptance','accepted_limitations':requirements.get('quality_warnings',[])}
    atomic_json(workspace.path/'acceptance.json',receipt)
    result.update(status='completed',answer='全部需求已由用户逐项验收，来源与交付文件核验通过。')
    result['completion']={'outcome':'completed',**artifacts,**requirements,
                          'all_requirements_verified':True,'semantic_review_required':False,'acceptance':receipt}
    atomic_json(workspace.path/'result.json',result)
    atomic_json(workspace.path/'delivery_check.json',result['completion'])
    workspace.log('delivery_accepted','',binding=binding,requirement_ids=requirement_ids)
    task=read_json(workspace.path/'task.json',{}).get('task','')
    # Only a human-accepted task enters cross-task memory; raw web instructions never do.
    content=('已验收任务：'+task+'\n结论摘要：'+(workspace.path/'research_report.md').read_text(encoding='utf-8').split('\n\n---',1)[0])[:1500]
    try:memory_id=memory_for(workspace).save(content,kind='accepted_task',task_id=workspace.task_id,metadata={'binding':binding})
    except ValueError:
        memory_id=None;workspace.log('memory_skipped','摘要包含受保护内容，未写入长期记忆')
    return {**result,'memory_id':memory_id}


def current_result(workspace):
    result=read_json(workspace.path/'result.json',{})
    if result.get('status')=='completed' and not (workspace.path/'.running').exists():
        accepted=read_json(workspace.path/'acceptance.json',{})
        artifacts=workspace.verify_delivery();requirements=evaluate_requirements(workspace)
        auto_valid=accepted.get('verification_mode')!='automatic_objective_checks' or completion_decision(workspace)['status']=='completed'
        if accepted.get('binding')!=review_binding(workspace) or not artifacts['artifacts_valid'] or not requirements['requirements_supported'] or not auto_valid:
            result['status']='needs_review' if artifacts['artifacts_valid'] and requirements['requirements_supported'] else 'partial'
            result['completion']={**artifacts,**requirements,'outcome':completion_outcome(result['status'],bool(workspace.sources())),'all_requirements_verified':False,'semantic_review_required':True,'acceptance_stale':True}
            result['answer']='验收后的文件或需求已变化，请重新核对并验收。'
            atomic_json(workspace.path/'result.json',result)
            memory_for(workspace).invalidate_task(workspace.task_id)
            workspace.log('acceptance_invalidated','交付版本已变化')
    return result


def request_delivery_revision(workspace,binding,note):
    if read_json(workspace.path/'demo_meta.json',{}).get('read_only'):raise ValueError('演示任务只读，不能修改验收')
    lock=workspace.path/'.running'
    try:lock.open('x').close()
    except FileExistsError as exc:raise ValueError('任务正在执行') from exc
    try:
        result=read_json(workspace.path/'result.json',{})
        if result.get('status')!='needs_review' or binding!=review_binding(workspace):raise ValueError('状态或交付版本已变化，请刷新后操作')
        if not 3<=len(note.strip())<=1000:raise ValueError('请用3–1000字符说明拒绝原因或修改要求')
        manifest=read_json(workspace.path/'task.json',{})
        manifest.update(review_mode='manual',revision_note=note.strip())
        atomic_json(workspace.path/'task.json',manifest)
        rows=read_json(workspace.path/'requirements.json',[])
        if len(rows)<16:
            define_requirements(workspace,[{'description':'用户修改要求：'+note.strip()[:580],'kind':'research'}],origin='user')
        (workspace.path/'acceptance.json').unlink(missing_ok=True)
        result.update(status='partial',answer='用户要求修改：'+note.strip())
        result['completion']={**workspace.verify_delivery(),**evaluate_requirements(workspace),'outcome':'partial','review_reason':note.strip()}
        atomic_json(workspace.path/'result.json',result)
        atomic_json(workspace.path/'delivery_check.json',result['completion'])
        workspace.log('delivery_revision_requested',note.strip())
        return result
    finally:lock.unlink(missing_ok=True)


def verify_existing_delivery(workspace):
    existing=current_result(workspace)
    if existing.get('status')=='completed':return existing
    lock=workspace.path/'.running'
    try:handle=lock.open('x')
    except FileExistsError as exc:raise ValueError('任务正在运行，不能重新核验') from exc
    handle.close()
    try:
        result=read_json(workspace.path/'result.json',{})
        if not result:raise ValueError('没有已保存的执行结果')
        artifacts=workspace.verify_delivery();requirements=evaluate_requirements(workspace)
        ready=artifacts['artifacts_valid'] and (not requirements['requirements'] or requirements['requirements_supported'])
        decision=completion_decision(workspace)
        result['status']=decision['status']
        if result['status']=='completed':automatic_receipt(workspace)
        result['answer']='已对保存的来源、文件与需求重新核验；未调用模型。'+decision['reason']
        result['completion']={'outcome':completion_outcome(result['status'],bool(workspace.sources())),**artifacts,**requirements}
        if result['status']=='completed':result['completion'].update(all_requirements_verified=True,semantic_review_required=False,acceptance=read_json(workspace.path/'acceptance.json',{}))
        atomic_json(workspace.path/'result.json',result);atomic_json(workspace.path/'delivery_check.json',result['completion'])
        workspace.log('delivery_rechecked','',additional_model_requests=0,ready=ready)
        return result
    finally:lock.unlink(missing_ok=True)


def rebind_unchanged_excerpts(workspace,old_sha,new_sha):
    """Carry exact, previously version-bound lines across deterministic re-export.

    Never marks pending/missing requirements satisfied and never changes source
    references, quoted text or human acceptance. Changed/removed lines stay stale.
    """
    if not old_sha or old_sha==new_sha:return []
    path=workspace.path/'research_report.md'
    if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest()!=new_sha:return []
    lines=path.read_text(encoding='utf-8').splitlines()
    rows=read_json(workspace.path/'requirements.json',[]);known={s['id'] for s in workspace.sources()};updated=[]
    for row in rows:
        result=row.get('result',{});parts=result.get('report_excerpts')
        if (result.get('status') not in {'satisfied','not_disclosed'}
                or result.get('report_sha256')!=old_sha
                or not isinstance(parts,list) or not parts or set(result.get('source_ids',[]))-known):continue
        if not all(isinstance(p,str) and p in lines for p in parts):continue
        result.update(report_sha256=new_sha,report_line_ids=sorted({lines.index(p)+1 for p in parts}))
        updated.append(row['id'])
    if updated:
        atomic_json(workspace.path/'requirements.json',rows)
        workspace.log('requirement_references_rebound','仅重绑定未改变的报告原文，未替代事实核验或用户验收',
                      requirement_ids=updated,old_sha256=old_sha,report_sha256=new_sha)
    return updated
