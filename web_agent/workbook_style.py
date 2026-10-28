"""Shared formatting for deterministic exported workbooks; no added data."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from datetime import datetime

def style_workbook(book):
    book.properties.title='WebTask 研究数据与核验证据'
    book.properties.creator='WebTask Agent'
    for index,ws in enumerate(book.worksheets,1):
        ws.sheet_view.showGridLines=False
        ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:1'
        ws.oddHeader.center.text=ws.title
        ws.oddFooter.right.text='Page &P / &N'
        for col in ws.columns:
            header=str(col[0].value or '')
            width=64 if header in {'原始摘录','缺失原因或说明','口径证据','说明'} else 48 if header=='来源网址' else 34 if header in {'对象','处理器'} else 19
            ws.column_dimensions[col[0].column_letter].width=width
            for c in col:
                c.font=Font(name='Microsoft YaHei',size=11,color='172B35')
                c.alignment=Alignment(vertical='center',wrap_text=True,horizontal='right' if type(c.value) in (int,float) else 'left',indent=1)
                if c.row>1:c.fill=PatternFill('solid',fgColor='F3F7F5' if c.row%2==0 else 'FFFFFF')
                if c.row>1 and header.endswith('查询日期') and isinstance(c.value,str):
                    try:c.value=datetime.strptime(c.value,'%Y-%m-%d');c.number_format='yyyy-mm-dd'
                    except ValueError:pass
                if type(c.value) in (int,float):c.number_format='#,##0' if float(c.value).is_integer() else '#,##0.00'
                c.border=Border(bottom=Side(style='hair',color='DDE6E2'))
                if c.value in ('未核实','官方未公布'):
                    c.fill=PatternFill('solid',fgColor='FFF0D9');c.font=Font(name='Microsoft YaHei',size=11,color='875417')
                if header=='来源网址' and isinstance(c.value,str) and c.value.startswith(('https://','http://')):
                    c.hyperlink=c.value;c.font=Font(name='Microsoft YaHei',size=10,color='087F70',underline='single')
        for row in ws.iter_rows(min_row=2):
            # Preserve full long quotations: row height follows wrapped text, within Excel's limit.
            lines=max((sum(max(1,(len(part)*1.5)//max(12,ws.column_dimensions[c.column_letter].width)+1) for part in str(c.value or '').split('\n')) for c in row),default=1)
            ws.row_dimensions[row[0].row].height=min(409,max(34,lines*16+10))
        for c in ws[1]:
            c.font=Font(name='Microsoft YaHei',size=11,bold=True,color='FFFFFF')
            c.fill=PatternFill('solid',fgColor='087F70')
            c.alignment=Alignment(vertical='center',horizontal='left',wrap_text=True)
        ws.row_dimensions[1].height=34
        if ws.max_row>1:
            table=Table(displayName=f'Delivery{index}',ref=ws.dimensions)
            table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True)
            ws.add_table(table)
            for text in ['未核实','官方未公布']:
                ws.conditional_formatting.add(ws.dimensions,CellIsRule(operator='equal',formula=['"'+text+'"'],
                    fill=PatternFill('solid',fgColor='FFF0D9'),font=Font(color='875417')))
