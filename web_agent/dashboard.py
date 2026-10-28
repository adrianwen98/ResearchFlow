"""Read-only view models for the local Gradio workspace. Never calls the LLM."""
from __future__ import annotations
import html
import json
import re
import time
from pathlib import Path
from urllib.parse import urlsplit
from .config import Settings
from .persistence import read_json, redact
from .completion import review_binding,current_result
from .workspace import Workspace

from .status import STATUS
from .completion import evaluate_requirements

NAMES={'update_plan':'制定与调整计划','load_skill':'加载任务方法','browser_navigate':'浏览网页',
       'browser_type':'输入网页搜索','browser_click':'点击页面控件','browser_select_option':'选择筛选条件',
       'browser_download':'通过浏览器下载','export_business_excel':'生成业务数据表',
       'browser_snapshot':'读取页面','browser_follow_link':'跟随页面链接','read_public_pdf':'读取 PDF',
       'download_public_file':'下载公开文件','inspect_downloaded_table':'检查表格字段',
       'aggregate_downloaded_table':'计算表格数据','record_source':'记录来源证据',
       'analyze_recorded_sources':'汇总来源','list_recorded_sources':'查看已存证据',
       'export_excel':'生成 Excel','export_report':'生成研究报告','request_human_review':'请求人工处理'}

def esc(value): return html.escape(str(value),quote=True)

def task_folder(task_id):
    if not re.fullmatch(r'[A-Za-z0-9_-]{1,56}',task_id or ''):
        raise ValueError('请输入有效的任务 ID。')
    return Settings.load().output_dir/task_id

def read_trace(folder):
    curated=folder/'public_trace.json'
    if curated.exists():return read_json(curated,[])[-1200:]
    path=folder/'trace.jsonl'
    if not path.exists(): return []
    rows=[]
    for line in path.read_text(encoding='utf-8',errors='replace').splitlines()[-1200:]:
        try: rows.append(json.loads(line))
        except json.JSONDecodeError: continue
    return rows

def snapshot(task_id='',result=None,started=None):
    if not task_id:
        return {'task_id':'','status':'idle','sources':[],'files':[],'metrics':{},'elapsed':0,'trace':[],
                'report':'','answer':'','plan':[],'pending':[],'task':'','requirements':[],'review_binding':'','requirement_checks':{},'artifacts_valid':False,'can_accept':False}
    folder=task_folder(task_id)
    saved=read_json(folder/'result.json',{})
    if saved.get('status')=='completed':saved=current_result(Workspace(folder.parent,task_id))
    if result is not None: saved=result.__dict__
    trace=read_trace(folder)
    sources=read_json(folder/'sources.json',[])
    metrics={**saved.get('metrics',{}),**read_json(folder/'usage.json',{})}
    metrics['tool_calls']=len(read_json(folder/'tool_journal.json',{}))
    metrics['errors']=sum(row['event']=='tool_error' for row in trace)
    metrics['limits']={**metrics.get('limits',{}),**read_json(folder/'run_limits.json',{})}
    meta=read_json(folder/'demo_meta.json',{})
    if 'actual_api_requests' in meta:metrics['requests']=meta['actual_api_requests']
    running=(folder/'.running').exists()
    status='running' if running else saved.get('status','running' if started else 'incomplete')
    report_path=folder/'research_report.md'
    report=report_path.read_text(encoding='utf-8') if report_path.exists() else ''
    files=[str(folder/name) for name in ('research_sources.xlsx','task_result.xlsx','research_report.md','trace.jsonl','sources.json') if (folder/name).is_file()]
    elapsed=round(time.monotonic()-started,1) if started else saved.get('elapsed_seconds',0)
    elapsed_label='当前任务的实际执行时间'
    if not started:
        initial_runs=read_json(folder.parent/'live_evaluation.json',[])
        initial_run=next((r for r in initial_runs if r.get('task_id')==task_id),None)
        if initial_run:
            elapsed=initial_run.get('elapsed_seconds',elapsed)
            elapsed_label='首次研究的实际执行时间'
    manifest=read_json(folder/'task.json',{})
    ws=Workspace(folder.parent,task_id)
    checks=evaluate_requirements(ws)
    artifacts=ws.verify_delivery() if report_path.exists() and (folder/'research_sources.xlsx').exists() else {'artifacts_valid':False}
    can_accept=status=='needs_review' and artifacts['artifacts_valid'] and checks['requirements_supported'] and not meta.get('read_only')
    return {'task_id':task_id,'status':status,'sources':sources,'files':files,'metrics':metrics,
            'elapsed':elapsed, 'elapsed_label':elapsed_label,
            'trace':trace,'report':report,'answer':saved.get('answer',''),'plan':read_json(folder/'plan.json',[]),
            'pending':saved.get('pending_review',[]),'task':manifest.get('task',''),
            'requirements':read_json(folder/'requirements.json',[]),'review_binding':review_binding(ws),
            'requirement_checks':checks,'artifacts_valid':artifacts['artifacts_valid'],'can_accept':can_accept}

def summary_html(data):
    label,tone=STATUS.get(data['status'],(data['status'],'neutral'));m=data['metrics']
    limit=m.get('limits',{}).get('model_requests');used=m.get('requests',0)
    budget=f"{used} / {limit}" if limit else f"{used} / 未记录"
    cards=[('任务状态',esc(label),data['task_id'] or '等待输入任务',tone),
           ('模型请求预算',budget,'累计已用 / 本任务上限',''),
           ('工具调用 / 错误',f"{m.get('tool_calls',0)} / {m.get('errors',0)}",'逻辑调用次数 / 实际工具错误',''),
           ('来源 / 本次执行',f"{len(data['sources'])} 条 · {float(data['elapsed']):.1f} 秒",data.get('elapsed_label','本次运行时长'), '')]
    return '<div class="metric-grid">'+''.join(f'<div class="metric-card"><span class="metric-label">{title}</span><div class="metric-value {tone}">{value}</div><span class="metric-foot">{esc(note)}</span></div>' for title,value,note,tone in cards)+'</div>'


def stage_html(data):
    status=data['status'];events={r['event'] for r in data['trace']}
    phase=3 if status in {'needs_review','completed'} else 2 if 'verification' in events else 1 if 'tool_start' in events else 0
    items=[]
    for i,name in enumerate(('规划任务','执行与调整','核对证据','人工验收')):
        state='done' if status=='completed' or (i<phase and data['task_id']) else 'current' if i==phase and data['task_id'] else ''
        if status=='needs_review' and i==3:name='等待人工验收'
        items.append(f'<div class="stage {state}"><span>{"✓" if state=="done" else f"0{i+1}"}</span>{name}</div>')
    return '<div class="stages">'+''.join(items)+'</div>'


def requirements_html(data):
    checks=data.get('requirement_checks',{}).get('requirements',[])
    if not checks:return '<p class="review-note">尚无验收需求。任务开始后由用户或Agent拆解。</p>'
    items=[]
    for r in checks:
        fact_status=r.get('fact_status','not_applicable')
        label=('引用校验通过 · 事实待核对' if r['passed'] and fact_status in {'review_required','not_independently_verified'}
               else '证据检查通过' if r['passed'] else '待补充' if r['status']=='pending' else '未通过')
        scope=r.get('source_quality_issues') or []
        # These fields share a cited source, but may not belong to this exact
        # requirement. Never claim a fuzzy field-to-requirement match.
        note=(('引用来源还包含待核验字段：'+'；'.join(scope[:4])+'。请核对是否与本项要求有关。') if scope else
              r.get('reason','') or ('来源与报告引用已核对；事实内容仍需人工或客观规则验证。' if r['passed'] and fact_status!='not_applicable'
                                     else '请核对事实与任务含义，程序检查不代替人工判断。' if r['passed']
                                     else '缺少有效来源、报告摘录、计算或交付文件。'))
        items.append(f'<div class="requirement-row"><strong>{esc(r["id"])} · {esc(label)}</strong><p>{esc(r["description"])}</p><small>{esc(note)}</small></div>')
    return '<p class="review-note">只有全部需求和文件检查通过，且任务处于交付待复核状态，才能确认完成。</p>'+''.join(items)


def result_summary_html(data):
    if not data['task_id']:return '<p class="review-note">打开真实演示，或创建一个新任务。</p>'
    passed=sum(r['passed'] for r in data.get('requirement_checks',{}).get('requirements',[]))
    total=len(data.get('requirements',[]))
    label=STATUS.get(data['status'],(data['status'],''))[0]
    method=read_json(task_folder(data['task_id'])/'demo_meta.json',{}).get('execution_label','')
    return f'<div class="result-summary"><strong>{esc(label)} · 需求检查 {passed}/{total}</strong><p>{esc(method)}</p><p>已记录 {len(data["sources"])} 条来源，生成 {len([p for p in data["files"] if p.endswith((".xlsx",".md"))])} 份报告/表格。{"仍需用户最终核对。" if data["status"]!="completed" else "用户验收已记录。"}</p></div>'


def workbook_preview_html(data):
    if not data['task_id']:return '<p class="review-note">业务Excel生成后在此预览。</p>'
    path=task_folder(data['task_id'])/'task_result.xlsx'
    if not path.is_file() or path.is_symlink():return '<p class="review-note">此任务尚未生成业务Excel；来源清单不等于业务数据表。</p>'
    try:
        import zipfile
        from openpyxl import load_workbook
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist())>64*1024*1024:raise ValueError('工作簿超过预览限制')
        book=load_workbook(path,read_only=True,data_only=True);tables=[]
        try:
            for ws in list(book.worksheets)[:5]:
                if ws.title in {'Provenance','Metadata'}:continue
                rows=list(ws.iter_rows(max_row=min(ws.max_row or 1,9),max_col=min(ws.max_column or 1,30),values_only=True))
                if not rows:continue
                head=''.join(f'<th>{esc(v if v is not None else "")}</th>' for v in rows[0])
                body=''.join('<tr>'+''.join(f'<td>{esc(v if v is not None else "")}</td>' for v in row)+'</tr>' for row in rows[1:])
                tables.append(f'<h3>{esc(ws.title)}</h3><p class="review-note">预览前8行 · 共{max(0,(ws.max_row or 1)-1)}行，完整内容见交付文件。</p><div class="table-scroll"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>')
        finally:book.close()
        return ''.join(tables)
    except Exception:return '<p class="review-note">当前工作簿无法安全预览，请检查文件核验结果。</p>'


def sources_html(data):
    rows=[]
    for item in data['sources']:
        url=item.get('url','')
        link=f'<a href="{esc(url)}" target="_blank" rel="noopener noreferrer">查看原始来源 ↗</a>' if urlsplit(url).scheme in {'http','https'} else '<span>链接不可用</span>'
        rows.append(f'''<article class="source-card"><div class="source-top"><span class="source-number">{esc(item.get('id',''))}</span><span class="source-domain">{esc(urlsplit(url).hostname or '')}</span><span class="mini-badge">已匹配观察记录</span></div><h3>{esc(item.get('title',''))}</h3><p>{esc(item.get('fact',''))}</p><blockquote>{esc(item.get('quote',''))}</blockquote>{link}</article>''')
    return '<div class="source-list">'+''.join(rows)+'</div>' if rows else empty_html('来源证据会出现在这里','实际读取网页或文件并记录来源后，可在此核对原文与链接。','source')

def empty_html(title='你的研究成果，从这里开始',description='在左侧描述任务。Agent 会规划步骤、调用工具，并把报告与来源整理在这里。',icon='report'):
    glyph={'report':'▤','source':'↗','trace':'≋','files':'↓'}[icon]
    return f'<div class="empty-state"><div class="empty-icon">{glyph}</div><h3>{esc(title)}</h3><p>{esc(description)}</p><div class="empty-tags"><span>可追溯来源</span><span>真实执行记录</span><span>文件交付</span></div></div>'

def trace_html(data):
    records={}; order=[]
    for row in data['trace']:
        key=row.get('call_id')
        if row['event']=='tool_start':
            records[key]={'name':row['detail'],'arguments':row.get('arguments',{}),'state':'running'}; order.append(key)
        elif key in records and row['event']=='tool_error': records[key]['error']=row.get('detail','')
        elif key in records and row['event']=='tool_end':
            records[key].update({'state':'error' if row.get('output_preview','').startswith(('ERROR:','下载失败')) else 'done','elapsed':row.get('elapsed_seconds',0),'output':row.get('output_preview','')})
    plans=''.join(f'<li>{esc(step)}</li>' for step in data['plan'])
    plan=f'<div class="plan-panel"><span class="section-eyebrow">当前执行计划</span><ol>{plans}</ol></div>' if plans else ''
    adjustments=[r for r in data['trace'] if r['event'] in {'plan_updated','task_resume','task_interrupted','tool_retry'}]
    plan+='<details class=adjustments><summary>调整、重试与恢复记录（'+str(len(adjustments))+'）</summary>'+''.join('<p>'+esc(r['time'][11:19]+' · '+r['event']+' · '+str(r.get('detail',''))[:350])+'</p>' for r in adjustments[-12:])+'</details>' if adjustments else ''
    items=[]
    for key in order[-50:]:
        rec=records[key]; args=rec['arguments']; info=args.get('url') or args.get('filename') or args.get('title') or ''
        name=NAMES.get(rec['name'],rec['name']); error=esc(redact(rec.get('error',''))[:220])
        items.append(f'<div class="trace-item {rec["state"]}"><span class="trace-mark">{"!" if rec["state"]=="error" else "✓" if rec["state"]=="done" else "·"}</span><div><div class="trace-title">{esc(name)}<span>{rec.get("elapsed",0):.2f}s</span></div><p>{esc(info) or esc(rec["name"])}</p>{f"<div class=trace-error>{error}</div>" if error else ""}<details class=tool-result><summary>参数与实际结果</summary><pre>{esc(redact(json.dumps(args,ensure_ascii=False)))}</pre><pre>{esc(redact(rec.get("output","等待结果")))}</pre></details></div></div>')
    if not items: return plan+empty_html('执行过程，逐步可见','任务开始后，这里会显示真实工具调用、错误与调整路径。','trace')
    return plan+'<div class="trace-list">'+''.join(items)+'</div>'

def history_choices():
    root=Settings.load().output_dir
    rows=[]
    if root.exists():
        for folder in root.iterdir():
            path=folder/'task.json'
            if path.is_file() and re.fullmatch(r'[A-Za-z0-9_-]{1,56}',folder.name):
                from .task_history import metadata
                meta=metadata(root,folder.name)
                if meta.get('deleted') or meta.get('archived'):continue
                item=read_json(path,{})
                title=re.sub(r'https?://\S+','',meta.get('title') or item.get('task','')).strip(' ，。、')[:34]
                rows.append((path.stat().st_mtime,(f'{folder.name} · {title}',folder.name)))
    return [row for _,row in sorted(rows,reverse=True)][:30]


def recommended_demo():
    choices={value for _,value in history_choices()}
    for candidate in ('interaction_end_to_end','penguins_public_demo','completion_research'):
        if candidate in choices:return candidate
    return next(iter(choices),'')
